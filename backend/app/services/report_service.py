"""
Report Export Service
- Generate PDF reports using ReportLab
- Generate DOCX reports using python-docx
"""
import os
import io
from datetime import datetime
from typing import Dict, List, Optional
from loguru import logger


def generate_pdf_report(paper_data: Dict, report_type: str = "full") -> bytes:
    """Generate a PDF report for a paper analysis."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch, cm
        from reportlab.lib import colors
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
            HRFlowable, PageBreak
        )
        from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=2*cm,
            leftMargin=2*cm,
            topMargin=2*cm,
            bottomMargin=2*cm,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Title'],
            fontSize=20,
            spaceAfter=12,
            textColor=colors.HexColor('#1e3a5f'),
        )
        heading_style = ParagraphStyle(
            'Heading',
            parent=styles['Heading2'],
            fontSize=13,
            spaceBefore=14,
            spaceAfter=6,
            textColor=colors.HexColor('#2563eb'),
        )
        body_style = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontSize=10,
            spaceAfter=6,
            leading=14,
            alignment=TA_JUSTIFY,
        )
        meta_style = ParagraphStyle(
            'Meta',
            parent=styles['Normal'],
            fontSize=9,
            textColor=colors.gray,
        )

        story = []

        # Header
        story.append(Paragraph("ResearchMind AI", meta_style))
        story.append(Paragraph("Analysis Report", title_style))
        story.append(Paragraph(f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}", meta_style))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#2563eb')))
        story.append(Spacer(1, 0.3*inch))

        # Paper title & metadata
        title = paper_data.get("title") or "Untitled Paper"
        story.append(Paragraph(title, title_style))

        authors = paper_data.get("authors")
        if authors:
            story.append(Paragraph(f"Authors: {', '.join(authors)}", meta_style))

        page_count = paper_data.get("page_count")
        if page_count:
            story.append(Paragraph(f"Pages: {page_count}", meta_style))

        story.append(Spacer(1, 0.2*inch))

        # Abstract
        abstract = paper_data.get("abstract")
        if abstract:
            story.append(Paragraph("Abstract", heading_style))
            story.append(Paragraph(abstract[:1500], body_style))

        # Summary
        summary = paper_data.get("summary") or {}
        if summary:
            story.append(Paragraph("Summary", heading_style))
            detailed = summary.get("detailed") or summary.get("brief", "")
            if detailed:
                story.append(Paragraph(detailed, body_style))

            section_sums = summary.get("section_summaries", {})
            for section, text in section_sums.items():
                if text:
                    story.append(Paragraph(section.replace("_", " ").title(), heading_style))
                    story.append(Paragraph(text, body_style))

        # Keywords
        keywords = paper_data.get("keywords") or []
        if keywords:
            story.append(Paragraph("Key Terms", heading_style))
            kw_text = " | ".join(k["keyword"] for k in keywords[:20] if isinstance(k, dict))
            story.append(Paragraph(kw_text, body_style))

        # Entities
        entities = paper_data.get("entities") or []
        if entities:
            story.append(Paragraph("Named Entities", heading_style))
            ent_rows = [["Entity", "Type"]]
            for ent in entities[:20]:
                ent_rows.append([ent.get("text", ""), ent.get("label", "")])
            t = Table(ent_rows, colWidths=[3.5*inch, 2*inch])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563eb')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f0f4ff')]),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.lightgrey),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ]))
            story.append(t)

        story.append(Spacer(1, 0.3*inch))
        story.append(HRFlowable(width="100%", thickness=0.5, color=colors.lightgrey))
        story.append(Paragraph("Generated by ResearchMind AI — AI Research Paper Intelligence System", meta_style))

        doc.build(story)
        return buffer.getvalue()

    except Exception as e:
        logger.error(f"PDF generation failed: {e}")
        raise


def generate_docx_report(paper_data: Dict) -> bytes:
    """Generate a DOCX report for a paper analysis."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # Title
        title_para = doc.add_heading("ResearchMind AI — Analysis Report", 0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph(f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}")
        doc.add_paragraph()

        # Paper metadata
        doc.add_heading(paper_data.get("title") or "Untitled Paper", 1)

        authors = paper_data.get("authors")
        if authors:
            doc.add_paragraph(f"Authors: {', '.join(authors)}")

        # Abstract
        abstract = paper_data.get("abstract")
        if abstract:
            doc.add_heading("Abstract", 2)
            doc.add_paragraph(abstract[:1500])

        # Summary
        summary = paper_data.get("summary") or {}
        if summary:
            doc.add_heading("Summary", 2)
            detailed = summary.get("detailed") or summary.get("brief", "")
            if detailed:
                doc.add_paragraph(detailed)

            section_sums = summary.get("section_summaries", {})
            for section, text in section_sums.items():
                if text:
                    doc.add_heading(section.replace("_", " ").title(), 3)
                    doc.add_paragraph(text)

        # Keywords
        keywords = paper_data.get("keywords") or []
        if keywords:
            doc.add_heading("Key Terms", 2)
            kw_text = ", ".join(k["keyword"] for k in keywords[:20] if isinstance(k, dict))
            doc.add_paragraph(kw_text)

        # Footer
        doc.add_paragraph()
        doc.add_paragraph("Generated by ResearchMind AI — AI Research Paper Intelligence System")

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    except Exception as e:
        logger.error(f"DOCX generation failed: {e}")
        raise
